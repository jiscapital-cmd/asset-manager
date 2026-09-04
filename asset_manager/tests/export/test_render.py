import io

from docx import Document
from pypdf import PdfReader

from asset_manager.export.render import markdown_to_docx_bytes, markdown_to_pdf_bytes


def test_markdown_to_docx_includes_title():
    content = "## Summary\nOccupancy is 94%."
    doc_bytes = markdown_to_docx_bytes("Champions Pointe — Sep 2026", content)
    doc = Document(io.BytesIO(doc_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Champions Pointe — Sep 2026" in all_text


def test_markdown_to_docx_renders_headings_and_body():
    content = "## Financial\nNOI came in under budget.\n## Risk\nModerate overall."
    doc_bytes = markdown_to_docx_bytes("Report", content)
    doc = Document(io.BytesIO(doc_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Financial" in all_text
    assert "NOI came in under budget." in all_text
    assert "Risk" in all_text
    assert "Moderate overall." in all_text


def test_markdown_to_docx_renders_bullet_lists():
    content = "## Findings\n- Occupancy dropped 3 points\n- Roof reserve is underfunded"
    doc_bytes = markdown_to_docx_bytes("Report", content)
    doc = Document(io.BytesIO(doc_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Occupancy dropped 3 points" in all_text
    assert "Roof reserve is underfunded" in all_text


def test_markdown_to_pdf_produces_valid_pdf_bytes():
    content = "## Summary\nOccupancy is 94%."
    pdf_bytes = markdown_to_pdf_bytes("Champions Pointe", content)
    assert pdf_bytes.startswith(b"%PDF")


def test_markdown_to_pdf_is_readable_and_has_at_least_one_page():
    content = "## Summary\nOccupancy is 94%."
    pdf_bytes = markdown_to_pdf_bytes("Champions Pointe", content)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1


def test_markdown_to_pdf_includes_title_text():
    content = "## Summary\nOccupancy is 94%."
    pdf_bytes = markdown_to_pdf_bytes("Champions Pointe Review", content)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    extracted = reader.pages[0].extract_text()
    assert "Champions Pointe Review" in extracted
