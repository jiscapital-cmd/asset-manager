import io

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from asset_manager.ingestion.loaders import (
    load_document,
    load_docx_text,
    load_pdf_text,
    load_text_file,
    load_xlsx_text,
)


def _make_pdf_bytes(pages: list[str]) -> bytes:
    writer = PdfWriter()
    for page_text in pages:
        writer.add_blank_page(width=200, height=200)
    # pypdf's add_blank_page doesn't support text injection directly;
    # for a real-text fixture we instead assert page count and let
    # load_pdf_text handle whatever text pypdf extracts (empty for blank
    # pages is fine — this test only verifies page-count behavior).
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_load_pdf_text_returns_one_entry_per_page():
    content = _make_pdf_bytes(["page one", "page two"])
    pages = load_pdf_text(content)
    assert [p for p, _ in pages] == [1, 2]


def test_load_docx_text_extracts_paragraphs():
    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    text = load_docx_text(buf.getvalue())
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_load_text_file_decodes_utf8():
    content = "rent roll data\nunit 1: $1200".encode("utf-8")
    assert load_text_file(content) == "rent roll data\nunit 1: $1200"


def test_load_document_dispatches_by_extension():
    txt_content = b"hello"
    result = load_document(txt_content, "notes.txt")
    assert result == [(1, "hello")]

    csv_content = b"a,b\n1,2"
    result = load_document(csv_content, "data.csv")
    assert result == [(1, "a,b\n1,2")]


def test_load_document_unknown_extension_raises():
    import pytest

    with pytest.raises(ValueError, match="Unsupported file type"):
        load_document(b"data", "file.xyz")


def _make_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append(["Unit", "Tenant", "Rent"])
    ws.append(["101", "Jane Doe", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_load_xlsx_text_includes_sheet_name_and_cell_values():
    text = load_xlsx_text(_make_xlsx_bytes())
    assert "Rent Roll" in text
    assert "Unit" in text
    assert "Jane Doe" in text
    assert "1200" in text


def test_load_document_dispatches_xlsx():
    result = load_document(_make_xlsx_bytes(), "rent_roll.xlsx")
    assert len(result) == 1
    page_number, text = result[0]
    assert page_number == 1
    assert "Jane Doe" in text
