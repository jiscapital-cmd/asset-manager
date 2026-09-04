"""Extract text from source documents, by file type."""

import io

from docx import Document
from pypdf import PdfReader


def load_pdf_text(content: bytes) -> list[tuple[int, str]]:
    """Return [(page_number, text), ...], 1-indexed."""
    reader = PdfReader(io.BytesIO(content))
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def load_docx_text(content: bytes) -> str:
    """Return all paragraph text joined with newlines."""
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


def load_text_file(content: bytes) -> str:
    """Decode a plain-text or CSV file as UTF-8."""
    return content.decode("utf-8")


def load_xlsx_text(content: bytes) -> str:
    """Flatten every sheet's cells into text, prefixed by sheet name — rent
    rolls and similar property-management exports are frequently .xlsx."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(", ".join(cells))
    return "\n".join(lines)


def load_document(content: bytes, filename: str) -> list[tuple[int, str]]:
    """Dispatch to the right loader by file extension.

    Returns a list of (page_or_row_number, text) — PDFs get one entry per
    page; everything else is treated as a single "page" numbered 1.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return load_pdf_text(content)
    if lower.endswith(".docx"):
        return [(1, load_docx_text(content))]
    if lower.endswith((".txt", ".csv", ".md")):
        return [(1, load_text_file(content))]
    if lower.endswith(".xlsx"):
        return [(1, load_xlsx_text(content))]
    raise ValueError(f"Unsupported file type: {filename}")
