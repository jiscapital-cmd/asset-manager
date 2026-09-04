"""Render a markdown report (as produced by the orchestrator) into DOCX/PDF.

Supports a deliberately small markdown subset — the orchestrator is
instructed (spec Section 3) to write reports as headings, paragraphs, and
bullet lists, which is all real asset-management reports need.
"""

import io


def _parse_markdown_lines(markdown_content: str) -> list[tuple[str, str]]:
    """Return [(kind, text), ...] where kind is 'heading', 'bullet', or 'paragraph'."""
    lines = []
    for raw_line in markdown_content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            lines.append(("heading", line[3:]))
        elif line.startswith("# "):
            lines.append(("heading", line[2:]))
        elif line.startswith("- "):
            lines.append(("bullet", line[2:]))
        else:
            lines.append(("paragraph", line))
    return lines


def markdown_to_docx_bytes(title: str, markdown_content: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    for kind, text in _parse_markdown_lines(markdown_content):
        if kind == "heading":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def markdown_to_pdf_bytes(title: str, markdown_content: str) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER)

    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    bullet_buffer: list[str] = []

    def flush_bullets():
        if bullet_buffer:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(b, styles["Normal"])) for b in bullet_buffer],
                    bulletType="bullet",
                )
            )
            bullet_buffer.clear()

    for kind, text in _parse_markdown_lines(markdown_content):
        if kind == "bullet":
            bullet_buffer.append(text)
            continue
        flush_bullets()
        if kind == "heading":
            story.append(Paragraph(text, styles["Heading2"]))
        else:
            story.append(Paragraph(text, styles["Normal"]))
        story.append(Spacer(1, 6))
    flush_bullets()

    doc.build(story)
    return buf.getvalue()
